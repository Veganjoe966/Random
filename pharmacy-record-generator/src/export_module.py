"""
Export Module — produces Excel (.xlsx) and CSV exports of generated records.

Excel workbook features:
  - "Dispensing Records" sheet: full data, frozen header, auto-width columns
  - "Summary" sheet: KPIs, top-10 drugs, payment breakdown, prescriber table
  - "Validation Report" sheet (if DEA issues exist)
  - Professional table formatting with alternating row colors
  - Simulation disclaimer in every workbook
"""

import io
import logging
from collections import Counter
from typing import Any, Dict, List, Optional

import pandas as pd

logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side,
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.dataframe import dataframe_to_rows
    _OPENPYXL_AVAILABLE = True
except ImportError:
    _OPENPYXL_AVAILABLE = False
    logger.warning("openpyxl not installed — Excel export unavailable.")

# ---------------------------------------------------------------------------
# Color Palette
# ---------------------------------------------------------------------------

_HEADER_FILL   = "1F3864"   # dark navy
_HEADER_FONT   = "FFFFFF"   # white
_ALT_ROW_FILL  = "EFF3FB"   # light blue-grey
_ACCENT_FILL   = "2E75B6"   # medium blue (section headers)
_ACCENT_FONT   = "FFFFFF"
_BORDER_COLOR  = "BDD7EE"
_TITLE_FILL    = "C6EFCE"   # green (summary title)
_WARN_FILL     = "FFC7CE"   # red (warnings / invalid)
_WARN_FONT     = "9C0006"


# ---------------------------------------------------------------------------
# Helper: thin border
# ---------------------------------------------------------------------------

def _thin_border() -> "Border":
    side = Side(style="thin", color=_BORDER_COLOR)
    return Border(left=side, right=side, top=side, bottom=side)


# ---------------------------------------------------------------------------
# ExportModule
# ---------------------------------------------------------------------------

class ExportModule:
    """
    Converts a dispensing records DataFrame into Excel or CSV bytes.

    Parameters
    ----------
    df                : pd.DataFrame — the generated records
    pharmacy_config   : dict         — pharmacy metadata
    date_range_str    : str          — human-readable date range label
    validation_issues : list         — DEA validation issue dicts (may be empty)
    """

    def __init__(
        self,
        df: pd.DataFrame,
        pharmacy_config: Dict[str, Any],
        date_range_str: str,
        validation_issues: Optional[List[Dict[str, Any]]] = None,
    ) -> None:
        self.df               = df.copy()
        self.pharmacy         = pharmacy_config
        self.date_range_str   = date_range_str
        self.validation_issues = validation_issues or []

    # ------------------------------------------------------------------
    # CSV Export
    # ------------------------------------------------------------------

    def to_csv_bytes(self) -> bytes:
        """Return UTF-8 CSV bytes of the full dispensing records."""
        buf = io.StringIO()
        # Prepend disclaimer row
        buf.write(
            "*** SIMULATED DATA — FOR COMPLIANCE TRAINING AND AUDIT TESTING ONLY ***\n"
        )
        buf.write(f"Pharmacy: {self.pharmacy.get('pharmacy_name', '')},,"
                  f"Date Range: {self.date_range_str}\n\n")
        self.df.to_csv(buf, index=False)
        return buf.getvalue().encode("utf-8")

    # ------------------------------------------------------------------
    # Excel Export
    # ------------------------------------------------------------------

    def to_excel_bytes(self) -> bytes:
        """
        Return .xlsx bytes with:
          - Dispensing Records sheet (data + formatting)
          - Summary sheet (KPIs + charts-ready tables)
          - Validation Report sheet (if DEA issues exist)
        """
        if not _OPENPYXL_AVAILABLE:
            raise RuntimeError(
                "openpyxl is not installed. "
                "Install it with: pip install openpyxl"
            )

        wb = Workbook()
        wb.remove(wb.active)    # remove default sheet

        self._write_records_sheet(wb)
        self._write_summary_sheet(wb)

        if self.validation_issues:
            self._write_validation_sheet(wb)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------
    # Records Sheet
    # ------------------------------------------------------------------

    def _write_records_sheet(self, wb: "Workbook") -> None:
        ws = wb.create_sheet("Dispensing Records")

        # --- Title rows ---
        pharmacy_name = self.pharmacy.get("pharmacy_name", "Unknown Pharmacy")
        ws.append([f"SIMULATED DISPENSING RECORD — {pharmacy_name.upper()}"])
        ws.append([f"Date Range: {self.date_range_str}"])
        ws.append(["*** FOR COMPLIANCE TRAINING AND AUDIT TESTING ONLY ***"])
        ws.append([])   # blank row

        for row_idx in [1, 2, 3]:
            cell = ws.cell(row=row_idx, column=1)
            cell.font = Font(bold=True, size=11, color="9C0006")

        # --- Column headers ---
        headers = [str(c).replace("_", " ").title() for c in self.df.columns]
        header_row = 5
        ws.append(headers)

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill   = PatternFill("solid", fgColor=_HEADER_FILL)
            cell.font   = Font(bold=True, color=_HEADER_FONT)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _thin_border()

        ws.row_dimensions[header_row].height = 28

        # --- Data rows ---
        for row_num, row_data in enumerate(
            dataframe_to_rows(self.df, index=False, header=False), start=header_row + 1
        ):
            ws.append(row_data)
            fill_color = _ALT_ROW_FILL if row_num % 2 == 0 else "FFFFFF"
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.fill      = PatternFill("solid", fgColor=fill_color)
                cell.alignment = Alignment(vertical="center")
                cell.border    = _thin_border()

        # --- Freeze header ---
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

        # --- Auto-width columns ---
        self._auto_width(ws, headers, start_row=header_row)

        # --- Tab color ---
        ws.sheet_properties.tabColor = _ACCENT_FILL

    # ------------------------------------------------------------------
    # Summary Sheet
    # ------------------------------------------------------------------

    def _write_summary_sheet(self, wb: "Workbook") -> None:
        ws = wb.create_sheet("Summary")

        def _section_header(row, text):
            ws.cell(row=row, column=1, value=text).font = Font(
                bold=True, size=12, color=_ACCENT_FONT
            )
            for col in range(1, 4):
                ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=_ACCENT_FILL)
            ws.row_dimensions[row].height = 20

        def _kv(row, key, value):
            kc = ws.cell(row=row, column=1, value=key)
            vc = ws.cell(row=row, column=2, value=value)
            kc.font = Font(bold=True)
            kc.alignment = Alignment(horizontal="right")
            vc.alignment = Alignment(horizontal="left")

        df = self.df
        schedule_col = "controlled_schedule"
        drug_col     = "drug_name"
        payment_col  = "payment_type"
        plan_col     = "insurance_plan"

        total = len(df)
        n_cii    = len(df[df[schedule_col] == "Schedule II"])
        n_ciii   = len(df[df[schedule_col] == "Schedule III"])
        n_civ    = len(df[df[schedule_col] == "Schedule IV"])
        n_cv     = len(df[df[schedule_col] == "Schedule V"])
        n_ctrl   = n_cii + n_ciii + n_civ + n_cv
        n_nctrl  = len(df[df[schedule_col] == "Non-Controlled"])
        n_cash   = len(df[df[payment_col] == "Cash"])
        n_ins    = len(df[df[payment_col] == "Insurance"])

        r = 1

        # Title
        ws.cell(row=r, column=1, value="DISPENSING RECORD — SIMULATION SUMMARY").font = Font(
            bold=True, size=14, color="9C0006"
        )
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        r += 1

        ws.cell(row=r, column=1, value="*** SIMULATED DATA — FOR TRAINING ONLY ***").font = Font(
            bold=True, color="9C0006"
        )
        r += 2

        # Pharmacy details
        _section_header(r, "  Pharmacy Information"); r += 1
        _kv(r, "Pharmacy Name:",  self.pharmacy.get("pharmacy_name", "")); r += 1
        _kv(r, "NPI:",            self.pharmacy.get("npi", "")); r += 1
        _kv(r, "DEA Number:",     self.pharmacy.get("dea_number", "")); r += 1
        _kv(r, "Address:",        self.pharmacy.get("address", "")); r += 1
        _kv(r, "Date Range:",     self.date_range_str); r += 2

        # Volume KPIs
        _section_header(r, "  Volume Summary"); r += 1
        _kv(r, "Total Dispensing Records:", total); r += 1
        _kv(r, "Total Controlled Substances:", n_ctrl); r += 1
        _kv(r, "  Schedule II (CII):",  n_cii); r += 1
        _kv(r, "  Schedule III (CIII):", n_ciii); r += 1
        _kv(r, "  Schedule IV (CIV):",  n_civ); r += 1
        _kv(r, "  Schedule V (CV):",    n_cv); r += 1
        _kv(r, "Total Non-Controlled:", n_nctrl); r += 2

        # Payment breakdown
        _section_header(r, "  Payment Breakdown"); r += 1
        _kv(r, "Cash Fills:", n_cash); r += 1
        _kv(r, "Insurance Fills:", n_ins); r += 1
        cash_pct = round(n_cash / total * 100, 1) if total else 0
        _kv(r, "Cash Percentage:", f"{cash_pct}%"); r += 2

        # Top 10 dispensed medications
        _section_header(r, "  Top 10 Dispensed Medications"); r += 1
        hdr_row = r
        for col, hdr in enumerate(["Rank", "Drug Name", "Count", "% of Total"], start=1):
            c = ws.cell(row=hdr_row, column=col, value=hdr)
            c.font   = Font(bold=True, color=_HEADER_FONT)
            c.fill   = PatternFill("solid", fgColor=_HEADER_FILL)
            c.alignment = Alignment(horizontal="center")
        r += 1

        drug_counts = Counter(df[drug_col].tolist())
        for rank, (drug, cnt) in enumerate(drug_counts.most_common(10), start=1):
            pct = round(cnt / total * 100, 1) if total else 0
            row_fill = _ALT_ROW_FILL if rank % 2 == 0 else "FFFFFF"
            for col_idx, val in enumerate([rank, drug, cnt, f"{pct}%"], start=1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.fill = PatternFill("solid", fgColor=row_fill)
                c.border = _thin_border()
            r += 1
        r += 1

        # Prescriber distribution (top 20)
        _section_header(r, "  Prescriber Distribution (Top 20)"); r += 1
        presc_col = "prescriber_name"
        spec_col  = "prescriber_specialty"

        hdr_row = r
        for col, hdr in enumerate(["Prescriber", "Specialty", "Fills", "% of Total"], start=1):
            c = ws.cell(row=hdr_row, column=col, value=hdr)
            c.font   = Font(bold=True, color=_HEADER_FONT)
            c.fill   = PatternFill("solid", fgColor=_HEADER_FILL)
            c.alignment = Alignment(horizontal="center")
        r += 1

        presc_counts = Counter(df[presc_col].tolist())
        for rank, (name, cnt) in enumerate(presc_counts.most_common(20), start=1):
            spec = df[df[presc_col] == name][spec_col].iloc[0] if len(
                df[df[presc_col] == name]
            ) > 0 else ""
            pct = round(cnt / total * 100, 1) if total else 0
            row_fill = _ALT_ROW_FILL if rank % 2 == 0 else "FFFFFF"
            for col_idx, val in enumerate([name, spec, cnt, f"{pct}%"], start=1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.fill = PatternFill("solid", fgColor=row_fill)
                c.border = _thin_border()
            r += 1
        r += 1

        # Insurance plan breakdown
        _section_header(r, "  Insurance Plan Breakdown"); r += 1
        ins_df = df[df[payment_col] == "Insurance"]
        hdr_row = r
        for col, hdr in enumerate(["Insurance Plan", "Fills", "% of Insurance Fills"], start=1):
            c = ws.cell(row=hdr_row, column=col, value=hdr)
            c.font   = Font(bold=True, color=_HEADER_FONT)
            c.fill   = PatternFill("solid", fgColor=_HEADER_FILL)
            c.alignment = Alignment(horizontal="center")
        r += 1

        plan_counts = Counter(ins_df[plan_col].tolist())
        for rank, (plan, cnt) in enumerate(plan_counts.most_common(), start=1):
            pct = round(cnt / n_ins * 100, 1) if n_ins else 0
            row_fill = _ALT_ROW_FILL if rank % 2 == 0 else "FFFFFF"
            for col_idx, val in enumerate([plan, cnt, f"{pct}%"], start=1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.fill = PatternFill("solid", fgColor=row_fill)
                c.border = _thin_border()
            r += 1

        # Auto-width summary sheet
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

        ws.sheet_properties.tabColor = "70AD47"

    # ------------------------------------------------------------------
    # Validation Report Sheet
    # ------------------------------------------------------------------

    def _write_validation_sheet(self, wb: "Workbook") -> None:
        ws = wb.create_sheet("Validation Report")

        ws.cell(row=1, column=1, value="DEA Validation Report").font = Font(
            bold=True, size=13, color="9C0006"
        )
        ws.cell(row=2, column=1, value=(
            f"{len(self.validation_issues)} issue(s) detected in uploaded prescriber file."
        ))
        ws.append([])

        headers = ["Row", "Prescriber Name", "DEA Number", "Status"]
        hdr_row = 4
        ws.append(headers)
        for col_idx, hdr in enumerate(headers, start=1):
            c = ws.cell(row=hdr_row, column=col_idx)
            c.font   = Font(bold=True, color=_HEADER_FONT)
            c.fill   = PatternFill("solid", fgColor=_HEADER_FILL)
            c.alignment = Alignment(horizontal="center")

        for i, issue in enumerate(self.validation_issues, start=1):
            row_num = hdr_row + i
            row_fill = _WARN_FILL if "INVALID" in issue.get("status", "") else _ALT_ROW_FILL
            ws.cell(row=row_num, column=1, value=issue.get("row", "")).fill    = PatternFill("solid", fgColor=row_fill)
            ws.cell(row=row_num, column=2, value=issue.get("name", "")).fill   = PatternFill("solid", fgColor=row_fill)
            ws.cell(row=row_num, column=3, value=issue.get("dea", "")).fill    = PatternFill("solid", fgColor=row_fill)
            ws.cell(row=row_num, column=4, value=issue.get("status", "")).fill = PatternFill("solid", fgColor=row_fill)

            if "INVALID" in issue.get("status", ""):
                ws.cell(row=row_num, column=4).font = Font(color=_WARN_FONT, bold=True)

        for col in ws.columns:
            max_len = max(
                (len(str(cell.value)) for cell in col if cell.value), default=10
            )
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)

        ws.sheet_properties.tabColor = "FF0000"

    # ------------------------------------------------------------------
    # Auto-Width Helper
    # ------------------------------------------------------------------

    @staticmethod
    def _auto_width(ws, headers: List[str], start_row: int) -> None:
        """Set column widths based on content."""
        from openpyxl.utils import get_column_letter

        for col_idx, header in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            max_width  = len(header)

            # Sample up to 500 rows for performance
            sample_end = min(ws.max_row, start_row + 500)
            for row_idx in range(start_row + 1, sample_end + 1):
                cell_val = ws.cell(row=row_idx, column=col_idx).value
                if cell_val:
                    max_width = max(max_width, len(str(cell_val)))

            ws.column_dimensions[col_letter].width = min(max_width + 3, 50)
